"""
Sprint 6 — Day 40: Screener endpoints.

Endpoints:
  GET /api/v1/screener            — List all screener presets
  GET /api/v1/screener/{preset}   — Get companies for a specific preset
"""

from __future__ import annotations

from pathlib import Path
from typing import List

import openpyxl
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

router = APIRouter(prefix="/api/v1/screener", tags=["screener"])

SCREENER_XLSX = (
    Path(__file__).resolve().parent.parent.parent.parent.parent
    / "Sprint-03-Screener-Peer-Comparison-Engine"
    / "output"
    / "screener_output.xlsx"
)

PRESET_NAMES = [
    "Quality Compounder",
    "Value Pick",
    "Growth Accelerator",
    "Dividend Champion",
    "Debt-Free Blue Chip",
    "Turnaround Watch",
]


def _load_preset(preset: str) -> List[dict]:
    """Load a screener preset from the Excel file."""
    wb = openpyxl.load_workbook(str(SCREENER_XLSX), read_only=True)
    if preset not in wb.sheetnames:
        wb.close()
        raise HTTPException(status_code=404, detail=f"Preset not found: {preset!r}")
    ws = wb[preset]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    return [dict(zip(headers, row)) for row in rows[1:]]


class PresetSummary(BaseModel):
    """Screener preset summary."""

    name: str
    company_count: int


@router.get("/", response_model=List[PresetSummary])
def list_presets():
    """List all available screener presets with company counts."""
    results = []
    for name in PRESET_NAMES:
        data = _load_preset(name)
        results.append(PresetSummary(name=name, company_count=len(data)))
    return results


@router.get("/{preset}")
def get_preset(preset: str):
    """Return all companies matching a screener preset."""
    return _load_preset(preset)
